"""Helpers for the PPTX build: theme, shapes, tables, entrance animations, fit checks."""
import copy
from lxml import etree
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from PIL import ImageFont

# ---------- theme ----------
BG, INK, CORAL, TEAL, STEEL, GREY, AMBER = "FAFBFC", "15323F", "D1495B", "2A9D8F", "386FA4", "8FA3AD", "E9C46A"
CORAL_T, TEAL_T, STEEL_T, AMBER_T, LINE, PANEL = "F6DDE1", "D9EFEC", "DCE7F2", "FBF1D3", "DDE4E8", "EEF2F4"
MUTED = "5B7380"  # darker grey for small text (contrast on BG)
FONT, MONO = "Calibri", "Courier New"
W, H = 13.333, 7.5
MX = 0.6  # side margin

FONT_FILES = {
    False: "/home/ngthtrong/.local/share/fonts/ms-office/calibri.ttf",
    True: "/home/ngthtrong/.local/share/fonts/ms-office/calibrib.ttf",
}
_fc = {}
WARNINGS = []


def rgb(h):
    return RGBColor.from_string(h)


def _font(size, bold):
    k = (round(size * 4), bold)
    if k not in _fc:
        _fc[k] = ImageFont.truetype(FONT_FILES[bool(bold)], int(round(size * 4)))
    return _fc[k]


def text_width_in(s, size, bold=False, mono=False):
    if mono:
        return len(s) * size * 0.6 / 72
    return _font(size, bold).getlength(s) / 4 / 72


def wrap_lines(s, size, width_in, bold=False, mono=False):
    n = 0
    for para in s.split("\n"):
        words = para.split(" ")
        cur = ""
        lines = 1
        for w in words:
            t = (cur + " " + w).strip()
            if text_width_in(t, size, bold, mono) <= width_in or not cur:
                cur = t
            else:
                lines += 1
                cur = w
        n += lines
    return n


def runs_text(paras):
    return "\n".join("".join(r[0] for r in p["runs"]) for p in paras)


# ---------- paragraph spec ----------
def P(text="", size=14, bold=False, color=INK, align="l", mono=False, italic=False, space=0, bullet=False, runs=None):
    """A paragraph. `runs` = list of (text, {bold,color,size,mono,italic}) to mix styles."""
    if runs is None:
        runs = [(text, {})]
    return dict(runs=runs, size=size, bold=bold, color=color, align=align, mono=mono, italic=italic, space=space, bullet=bullet)


def B(text, **kw):
    kw.setdefault("bullet", True)
    kw.setdefault("space", 4)
    return P(text, **kw)


def _apply_paras(tf, paras, width_in):
    first = True
    lines_h = 0.0
    for p in paras:
        para = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        para.alignment = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}[p["align"]]
        if p["space"]:
            para.space_before = Pt(p["space"])
        indent = 0.0
        if p["bullet"]:
            pPr = para._p.get_or_add_pPr()
            pPr.set("marL", str(int(Inches(0.2))))
            pPr.set("indent", str(-int(Inches(0.2))))
            for tag in ("a:buNone", "a:buChar", "a:buAutoNum"):
                for e in pPr.findall(qn(tag)):
                    pPr.remove(e)
            clr = etree.SubElement(pPr, qn("a:buClr"))
            etree.SubElement(clr, qn("a:srgbClr")).set("val", CORAL)
            etree.SubElement(pPr, qn("a:buFont")).set("typeface", "Arial")
            etree.SubElement(pPr, qn("a:buChar")).set("char", "•")
            indent = 0.2
        full = ""
        for txt, st in p["runs"]:
            r = para.add_run()
            r.text = txt
            f = r.font
            f.name = MONO if st.get("mono", p["mono"]) else FONT
            f.size = Pt(st.get("size", p["size"]))
            f.bold = st.get("bold", p["bold"])
            f.italic = st.get("italic", p["italic"])
            f.color.rgb = rgb(st.get("color", p["color"]))
            full += txt
        bold_any = p["bold"] or any(st.get("bold") for _, st in p["runs"])
        n = wrap_lines(full, p["size"], width_in - indent, bold_any, p["mono"])
        lines_h += n * p["size"] * 1.2 / 72 + p["space"] / 72
    return lines_h


def _noshadow(shp):
    spPr = shp._element.spPr
    if spPr.find(qn("a:effectLst")) is None:
        etree.SubElement(spPr, qn("a:effectLst"))


def name(shp, nm):
    shp.name = nm
    return shp


class Deck:
    def __init__(self):
        self.prs = Presentation()
        self.prs.slide_width = Inches(W)
        self.prs.slide_height = Inches(H)
        self.cur = None
        self.n = 0
        self.anims = {}

    def slide(self, title, dark=False, section=0, notes="", size=26, title_h=0.95, footer=True, bg=None):
        s = self.prs.slides.add_slide(self.prs.slide_layouts[5])  # Title Only
        self.n += 1
        self.cur = s
        self.dark = dark
        self.anims[self.n] = []
        s.background.fill.solid()
        s.background.fill.fore_color.rgb = rgb(bg or (INK if dark else BG))
        t = s.shapes.title
        t.left, t.top, t.width, t.height = Inches(MX), Inches(0.42), Inches(W - 2 * MX), Inches(title_h)
        tf = t.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = MSO_ANCHOR.TOP
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        tf.text = ""
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.LEFT
        r = p.add_run()
        r.text = title
        r.font.name, r.font.size, r.font.bold = FONT, Pt(size), True
        r.font.color.rgb = rgb(BG if dark else INK)
        need = wrap_lines(title, size, W - 2 * MX, True) * size * 1.2 / 72
        if need > title_h + 0.01:
            WARNINGS.append(f"slide {self.n}: title needs {need:.2f}in > {title_h}")
        if notes:
            s.notes_slide.notes_text_frame.text = notes
        if footer:
            self._footer(section)
        return s

    def _footer(self, section):
        col = GREY if self.dark else MUTED
        self.text(MX, 7.08, 7.0, 0.25, [P("Residue-level prediction of degron mutation effects  ·  CT553H  ·  Nguyễn Thế Trọng", size=10, color=col)], nm="footer")
        # progress motif: ten residue cells, one per section
        x0 = W - MX - 10 * 0.22 - 0.9
        for i in range(10):
            on = i + 1 == section
            c = CORAL if on else (("2F4B59" if self.dark else LINE))
            self.rect(x0 + i * 0.22, 7.13, 0.16, 0.16, fill=c, nm=f"progress {i + 1}")
        self.text(W - MX - 0.75, 7.08, 0.75, 0.25, [P(f"{self.n}", size=10, color=col, align="r")], nm="slide number")

    # ---- primitives ----
    def text(self, x, y, w, h, paras, anchor="t", margin=0.0, nm=None, fill=None, line=None, check=True):
        s = self.cur
        tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = Inches(margin)
        tf.margin_top = tf.margin_bottom = Inches(margin * 0.6)
        tf.vertical_anchor = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE, "b": MSO_ANCHOR.BOTTOM}[anchor]
        if fill:
            tb.fill.solid()
            tb.fill.fore_color.rgb = rgb(fill)
        if line:
            tb.line.color.rgb = rgb(line)
        need = _apply_paras(tf, paras, w - 2 * margin) + margin * 1.2
        if check and need > h + 0.02:
            WARNINGS.append(f"slide {self.n}: text '{runs_text(paras)[:40]}' needs {need:.2f}in > {h:.2f}in")
        if nm:
            tb.name = nm
        tb._need = need
        return tb

    def rect(self, x, y, w, h, fill=None, line=None, lw=1.0, shape="rect", dash=False, nm=None, paras=None, anchor="m", margin=0.08, radius=None, transp=None):
        kind = {"rect": MSO_SHAPE.RECTANGLE, "round": MSO_SHAPE.ROUNDED_RECTANGLE, "oval": MSO_SHAPE.OVAL,
                "tri": MSO_SHAPE.ISOSCELES_TRIANGLE, "rarrow": MSO_SHAPE.RIGHT_ARROW, "darrow": MSO_SHAPE.DOWN_ARROW,
                "chev": MSO_SHAPE.CHEVRON, "can": MSO_SHAPE.CAN, "pent": MSO_SHAPE.PENTAGON, "star": MSO_SHAPE.FLOWCHART_SUMMING_JUNCTION}[shape]
        shp = self.cur.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
        _noshadow(shp)
        if shape == "round":
            shp.adjustments[0] = radius if radius is not None else min(0.5, 0.1 / max(min(w, h), 0.01))
        if fill:
            shp.fill.solid()
            shp.fill.fore_color.rgb = rgb(fill)
        else:
            shp.fill.background()
        if line:
            shp.line.color.rgb = rgb(line)
            shp.line.width = Pt(lw)
            if dash:
                ln = shp.line._get_or_add_ln()
                pd = etree.SubElement(ln, qn("a:prstDash"))
                pd.set("val", "dash")
        else:
            shp.line.fill.background()
        tf = shp.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = Inches(margin)
        tf.margin_top = tf.margin_bottom = Inches(margin * 0.5)
        tf.vertical_anchor = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE, "b": MSO_ANCHOR.BOTTOM}[anchor]
        if paras:
            need = _apply_paras(tf, paras, w - 2 * margin) + margin
            if need > h + 0.02:
                WARNINGS.append(f"slide {self.n}: shape text '{runs_text(paras)[:40]}' needs {need:.2f}in > {h:.2f}in")
        if nm:
            shp.name = nm
        return shp

    def line(self, x1, y1, x2, y2, color=GREY, lw=1.5, dash=False, arrow=False, nm=None):
        c = self.cur.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
        c.line.color.rgb = rgb(color)
        c.line.width = Pt(lw)
        ln = c.line._get_or_add_ln()
        if dash:
            etree.SubElement(ln, qn("a:prstDash")).set("val", "dash")
        if arrow:
            etree.SubElement(ln, qn("a:tailEnd")).set("type", "triangle")
        if nm:
            c.name = nm
        return c

    def pic(self, path, x, y, w=None, h=None, alt="", nm=None):
        kw = {}
        if w:
            kw["width"] = Inches(w)
        if h:
            kw["height"] = Inches(h)
        p = self.cur.shapes.add_picture(path, Inches(x), Inches(y), **kw)
        p._element.nvPicPr.cNvPr.set("descr", alt)
        if nm:
            p.name = nm
        return p

    def table(self, x, y, colw, rows, size=11.5, header=True, fills=None, bold_first_col=False, hdr_fill=INK, pad=0.06,
              min_row=0.3, nm=None, aligns=None, row_fills=None, mono_cols=()):
        """rows: list of lists of str or list-of-runs spec ([(text,{style})]). Returns (graphicFrame, bottom_y)."""
        s = self.cur
        nr, nc = len(rows), len(colw)

        def cell_runs(c):
            return c if isinstance(c, list) else [(c, {})]

        heights = []
        for ri, row in enumerate(rows):
            mh = min_row
            for ci, c in enumerate(row):
                runs = cell_runs(c)
                txt = "".join(t for t, _ in runs)
                bold = (header and ri == 0) or (bold_first_col and ci == 0) or any(st.get("bold") for _, st in runs)
                n = wrap_lines(txt, size, colw[ci] - 2 * pad - 0.02, bold, ci in mono_cols)
                mh = max(mh, n * size * 1.2 / 72 + 0.1)
            heights.append(mh)
        total = sum(heights)
        gf = s.shapes.add_table(nr, nc, Inches(x), Inches(y), Inches(sum(colw)), Inches(total))
        if nm:
            gf.name = nm
        tbl = gf.table
        tblPr = tbl._tbl.tblPr
        for a in ("firstRow", "bandRow"):
            tblPr.set(a, "0")
        sid = tblPr.find(qn("a:tableStyleId"))
        if sid is not None:
            sid.text = "{2D5ABB26-0587-4C30-8999-92F81FD0307C}"  # No Style, No Grid
        for ci, cw in enumerate(colw):
            tbl.columns[ci].width = Inches(cw)
        for ri, row in enumerate(rows):
            tbl.rows[ri].height = Inches(heights[ri])
            for ci, c in enumerate(row):
                cell = tbl.cell(ri, ci)
                cell.margin_left = cell.margin_right = Inches(pad)
                cell.margin_top = cell.margin_bottom = Inches(0.03)
                cell.vertical_anchor = MSO_ANCHOR.MIDDLE
                is_h = header and ri == 0
                fill = hdr_fill if is_h else (row_fills[ri] if row_fills and row_fills.get(ri) else (fills.get((ri, ci)) if fills else None))
                if fill is None and not is_h:
                    fill = "FFFFFF" if ri % 2 else "F4F7F8"
                cell.fill.solid()
                cell.fill.fore_color.rgb = rgb(fill)
                tf = cell.text_frame
                tf.word_wrap = True
                para = tf.paragraphs[0]
                al = (aligns[ci] if aligns else "l")
                para.alignment = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}[al]
                for txt, st in cell_runs(c):
                    r = para.add_run()
                    r.text = txt
                    f = r.font
                    f.name = MONO if (ci in mono_cols or st.get("mono")) else FONT
                    f.size = Pt(st.get("size", size))
                    f.bold = True if is_h else st.get("bold", bold_first_col and ci == 0)
                    f.italic = st.get("italic", False)
                    f.color.rgb = rgb(BG if is_h else st.get("color", INK))
                # thin bottom border
                tcPr = cell._tc.get_or_add_tcPr()
                for side in ("a:lnL", "a:lnR", "a:lnT"):
                    ln = etree.SubElement(tcPr, qn(side))
                    ln.set("w", "0")
                    etree.SubElement(ln, qn("a:noFill"))
                ln = etree.SubElement(tcPr, qn("a:lnB"))
                ln.set("w", "6350")
                sf = etree.SubElement(ln, qn("a:solidFill"))
                etree.SubElement(sf, qn("a:srgbClr")).set("val", LINE)
                # schema order: borders must precede fill in tcPr
                fillel = [e for e in tcPr if e.tag in (qn("a:solidFill"), qn("a:noFill")) and e.getparent() is tcPr]
                for e in fillel:
                    tcPr.remove(e)
                    tcPr.append(e)
        gf._bottom = y + total
        return gf, y + total

    def shift(self, dy):
        keep = lambda n: n == 'footer' or n.startswith('progress') or n == 'slide number'
        for sh in self.cur.shapes:
            if sh == self.cur.shapes.title or keep(sh.name):
                continue
            sh.top = sh.top - Inches(dy)

    # ---- animations ----
    def anim(self, *clicks):
        """Each click = list of shapes (or (shape, delay_ms) tuples) that fade in together.
        Shapes listed here should be present on the slide; they are hidden until their click."""
        self.anims[self.n] = [list(c) for c in clicks]

    def finish(self):
        for idx, slide in enumerate(self.prs.slides, start=1):
            _transition(slide)
            clicks = self.anims.get(idx) or []
            if clicks:
                _timing(slide, clicks)


P_NS = "http://schemas.openxmlformats.org/presentationml/2006/main"


def _transition(slide):
    sld = slide._element
    if sld.find(qn("p:transition")) is not None:
        return
    tr = etree.Element(qn("p:transition"))
    tr.set("spd", "med")
    etree.SubElement(tr, qn("p:fade"))
    clr = sld.find(qn("p:clrMapOvr"))
    if clr is not None:
        clr.addnext(tr)
    else:
        sld.find(qn("p:cSld")).addnext(tr)


def _timing(slide, clicks):
    ids = iter(range(3, 10000))
    pars = []
    for click in clicks:
        inner = []
        for k, item in enumerate(click):
            shp, delay = (item if isinstance(item, tuple) else (item, 0))
            spid = shp.shape_id
            node = "clickEffect" if k == 0 else "withEffect"
            inner.append(
                f'<p:par><p:cTn id="{next(ids)}" presetID="10" presetClass="entr" presetSubtype="0" fill="hold" nodeType="{node}">'
                f'<p:stCondLst><p:cond delay="{delay}"/></p:stCondLst><p:childTnLst>'
                f'<p:set><p:cBhvr><p:cTn id="{next(ids)}" dur="1" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst></p:cTn>'
                f'<p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl><p:attrNameLst><p:attrName>style.visibility</p:attrName></p:attrNameLst></p:cBhvr>'
                f'<p:to><p:strVal val="visible"/></p:to></p:set>'
                f'<p:animEffect transition="in" filter="fade"><p:cBhvr><p:cTn id="{next(ids)}" dur="500"/>'
                f'<p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl></p:cBhvr></p:animEffect>'
                f'</p:childTnLst></p:cTn></p:par>'
            )
        pars.append(
            f'<p:par><p:cTn id="{next(ids)}" fill="hold"><p:stCondLst><p:cond delay="indefinite"/></p:stCondLst><p:childTnLst>'
            f'<p:par><p:cTn id="{next(ids)}" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>'
            + "".join(inner) + '</p:childTnLst></p:cTn></p:par></p:childTnLst></p:cTn></p:par>'
        )
    xml = (
        f'<p:timing xmlns:p="{P_NS}"><p:tnLst><p:par><p:cTn id="1" dur="indefinite" restart="never" nodeType="tmRoot"><p:childTnLst>'
        f'<p:seq concurrent="1" nextAc="seek"><p:cTn id="2" dur="indefinite" nodeType="mainSeq"><p:childTnLst>'
        + "".join(pars) +
        '</p:childTnLst></p:cTn><p:prevCondLst><p:cond evt="onPrev" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:prevCondLst>'
        '<p:nextCondLst><p:cond evt="onNext" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:nextCondLst></p:seq>'
        '</p:childTnLst></p:cTn></p:par></p:tnLst></p:timing>'
    )
    sld = slide._element
    old = sld.find(qn("p:timing"))
    if old is not None:
        sld.remove(old)
    el = etree.fromstring(xml)
    tr = sld.find(qn("p:transition"))
    (tr if tr is not None else sld.find(qn("p:cSld"))).addnext(el)
