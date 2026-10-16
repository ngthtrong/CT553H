"""Rough layout preview of the PPTX (no LibreOffice available): draws shapes, text, tables and pictures with PIL.
Used for visual QA of overflow/overlap only; PowerPoint rendering will differ slightly."""
import sys, os
from pptx import Presentation
from pptx.util import Emu
from pptx.enum.shapes import MSO_SHAPE_TYPE, MSO_SHAPE
from PIL import Image, ImageDraw, ImageFont

PPI = 110
FONTS = {(0, 0): "/home/ngthtrong/.local/share/fonts/ms-office/calibri.ttf", (1, 0): "/home/ngthtrong/.local/share/fonts/ms-office/calibrib.ttf"}
MONO = "/usr/share/fonts/truetype/dejavu/DejaVuSansMono.ttf"
SYM = "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"
_c = {}


def font(sz, bold=False, mono=False):
    k = (round(sz * PPI / 72), bold, mono)
    if k not in _c:
        _c[k] = ImageFont.truetype(MONO if mono else FONTS[(int(bold), 0)], max(6, int(round(sz * PPI / 72))))
    return _c[k]


def hx(c):
    return tuple(int(str(c)[i:i + 2], 16) for i in (0, 2, 4))


def emu(v):
    return int(v / 914400 * PPI)


def draw_text(dr, tf, x, y, w, h, ml=0.0, mt=0.0, anchor="t"):
    blocks = []
    tot = 0
    for p in tf.paragraphs:
        runs = [(r.text, r.font) for r in p.runs]
        if not runs:
            blocks.append((0, [], p)); continue
        r0 = runs[0][1]
        sz = r0.size.pt if r0.size else 14
        bold = any(f.bold for _, f in runs)
        mono = (r0.name or "").startswith("Courier")
        fnt = font(sz, bold, mono)
        marl = int(p._p.pPr.get("marL", 0)) if p._p.pPr is not None and p._p.pPr.get("marL") else 0
        indent = emu(marl)
        words = "".join(t for t, _ in runs).split(" ")
        lines, cur = [], ""
        avail = w - 2 * ml - indent
        for wd in words:
            t = (cur + " " + wd).strip()
            if fnt.getlength(t) <= avail or not cur:
                cur = t
            else:
                lines.append(cur); cur = wd
        lines.append(cur)
        sb = p.space_before.pt * PPI / 72 if p.space_before else 0
        lh = sz * 1.2 * PPI / 72
        blocks.append((sb + lh * len(lines), lines, p, fnt, runs, indent, lh, sb, bold))
        tot += sb + lh * len(lines)
    cy = y + mt + ({"t": 0, "m": (h - 2 * mt - tot) / 2, "b": h - 2 * mt - tot}[anchor])
    for b in blocks:
        if len(b) == 3:
            continue
        _, lines, p, fnt, runs, indent, lh, sb, bold = b
        cy += sb
        col = runs[0][1].color.rgb if runs[0][1].color and runs[0][1].color.type else None
        col = hx(col) if col else (0, 0, 0)
        al = p.alignment
        for i, ln in enumerate(lines):
            tw = fnt.getlength(ln)
            if al is not None and int(al) == 2:
                tx = x + ml + indent + (w - 2 * ml - indent - tw) / 2
            elif al is not None and int(al) == 3:
                tx = x + w - ml - tw
            else:
                tx = x + ml + indent
            dr.text((tx, cy), ln, font=fnt, fill=col)
            if indent and i == 0:
                dr.text((x + ml, cy), "•", font=fnt, fill=(209, 73, 91))
            cy += lh
    return tot + 2 * mt


def render(pptx, outdir, pages=None):
    prs = Presentation(pptx)
    os.makedirs(outdir, exist_ok=True)
    for si, s in enumerate(prs.slides, 1):
        if pages and si not in pages:
            continue
        W, H = emu(prs.slide_width), emu(prs.slide_height)
        bgc = s.background.fill.fore_color.rgb
        im = Image.new("RGB", (W, H), hx(bgc))
        dr = ImageDraw.Draw(im)
        for sh in s.shapes:
            x, y, w, h = emu(sh.left), emu(sh.top), emu(sh.width), emu(sh.height)
            if sh.shape_type == MSO_SHAPE_TYPE.PICTURE:
                try:
                    pi = Image.open(sh.image.blob if False else __import__("io").BytesIO(sh.image.blob)).convert("RGB").resize((max(w, 1), max(h, 1)))
                    im.paste(pi, (x, y))
                except Exception:
                    dr.rectangle([x, y, x + w, y + h], outline=(200, 0, 0))
                continue
            if sh.has_table if hasattr(sh, "has_table") else False:
                tbl = sh.table
                cy = y
                for r in tbl.rows:
                    rh = emu(r.height)
                    cx = x
                    for ci, c in enumerate(r.cells):
                        cw = emu(tbl.columns[ci].width)
                        try:
                            fc = c.fill.fore_color.rgb
                            dr.rectangle([cx, cy, cx + cw, cy + rh], fill=hx(fc))
                        except Exception:
                            pass
                        need = draw_text(dr, c.text_frame, cx, cy, cw, rh, ml=emu(c.margin_left), mt=emu(c.margin_top), anchor="m")
                        if need > rh + 3:
                            dr.rectangle([cx, cy, cx + cw, cy + rh], outline=(255, 0, 0), width=2)
                        dr.line([cx, cy + rh, cx + cw, cy + rh], fill=(221, 228, 232))
                        cx += cw
                    cy += rh
                continue
            if sh.shape_type == MSO_SHAPE_TYPE.LINE or sh.__class__.__name__ == "Connector":
                try:
                    col = hx(sh.line.color.rgb)
                except Exception:
                    col = (140, 160, 170)
                dr.line([emu(sh.begin_x), emu(sh.begin_y), emu(sh.end_x), emu(sh.end_y)], fill=col, width=max(1, int(sh.line.width.pt * PPI / 72)) if sh.line.width else 2)
                continue
            fill = None
            try:
                if sh.fill.type == 1:
                    fill = hx(sh.fill.fore_color.rgb)
            except Exception:
                pass
            line = None
            try:
                if sh.line.fill.type == 1:
                    line = hx(sh.line.color.rgb)
            except Exception:
                pass
            lw = max(1, int((sh.line.width.pt if sh.line.width else 1) * PPI / 72))
            ast = None
            try:
                ast = sh.auto_shape_type
            except Exception:
                pass
            if ast == MSO_SHAPE.OVAL:
                dr.ellipse([x, y, x + w, y + h], fill=fill, outline=line, width=lw)
            elif ast == MSO_SHAPE.ROUNDED_RECTANGLE:
                dr.rounded_rectangle([x, y, x + w, y + h], radius=min(w, h) * 0.18, fill=fill, outline=line, width=lw)
            elif ast == MSO_SHAPE.RIGHT_ARROW:
                dr.polygon([(x, y + h * .25), (x + w * .6, y + h * .25), (x + w * .6, y), (x + w, y + h / 2), (x + w * .6, y + h), (x + w * .6, y + h * .75), (x, y + h * .75)], fill=fill, outline=line)
            elif ast == MSO_SHAPE.DOWN_ARROW:
                dr.polygon([(x + w * .25, y), (x + w * .75, y), (x + w * .75, y + h * .6), (x + w, y + h * .6), (x + w / 2, y + h), (x, y + h * .6), (x + w * .25, y + h * .6)], fill=fill, outline=line)
            elif ast == MSO_SHAPE.ISOSCELES_TRIANGLE:
                dr.polygon([(x + w / 2, y), (x + w, y + h), (x, y + h)], fill=fill, outline=line)
            elif ast == MSO_SHAPE.CAN:
                dr.rounded_rectangle([x, y, x + w, y + h], radius=min(w, h) * 0.25, fill=fill, outline=line, width=lw)
            elif fill or line:
                dr.rectangle([x, y, x + w, y + h], fill=fill, outline=line, width=lw)
            if sh.has_text_frame and sh.text_frame.text.strip():
                tf = sh.text_frame
                anc = {1: "t", 3: "m", 4: "b"}.get(int(tf.vertical_anchor) if tf.vertical_anchor is not None else 1, "t")
                need = draw_text(dr, tf, x, y, w, h, ml=emu(tf.margin_left), mt=emu(tf.margin_top), anchor=anc)
                if need > h + 4:
                    dr.rectangle([x, y, x + w, y + h], outline=(255, 0, 0), width=2)
        im.save(os.path.join(outdir, f"slide-{si:02d}.png"))
    print("rendered", outdir)


if __name__ == "__main__":
    pages = [int(a) for a in sys.argv[3:]] or None
    render(sys.argv[1], sys.argv[2], pages)
